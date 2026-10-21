"""把 排課結果.json 輸出成：課表檢視網頁（HTML）與灰階列印用 Excel。

用法：python engine/export.py 試跑/排課結果.json
"""
import json
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "web" / "課表檢視.template.html"
ROWS = [1, 2, "大下課", 3, 4, "午餐・午休", 5, 6, 7]


def insights(r):
    """把沒做到的「盡量」規則，整理成白話提醒：問題集中在誰、可以怎麼調。"""
    out = []
    for rule in r["rules"]:
        if rule["ok"] is not False:
            continue
        def who(x):
            # 「1年1班 體育（體育甲）…」以班級開頭 → 括號內是老師；其他訊息以老師開頭
            first = x.split(" ")[0]
            m = re.search(r"（([^）]+)）", x)
            return m.group(1) if re.fullmatch(r"\d年\d+班", first) and m else first
        top = Counter(who(x) for x in rule["details"]).most_common(2)
        focus = "、".join(f"{t}（{n}）" for t, n in top)
        out.append(dict(rule=rule["id"], text=f"{rule['id']}「{rule['content']}」有 {rule['n']} 處沒做到，主要集中在：{focus}。"))
    return out


def write_html(r, out):
    html = TEMPLATE.read_text(encoding="utf-8")
    payload = json.dumps(r, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    out.write_text(html.replace("__DATA__", payload), encoding="utf-8")


# ---------- Excel（灰階） ----------
THIN = Side(style="thin", color="808080")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FILL = {"head": PatternFill("solid", fgColor="D9D9D9"), "lock": PatternFill("solid", fgColor="BFBFBF"),
        "off": PatternFill("solid", fgColor="808080"), "lang": PatternFill("solid", fgColor="EDEDED"),
        "sped": PatternFill("solid", fgColor="BFBFBF"), "keep": PatternFill("solid", fgColor="D9D9D9"), "must": PatternFill("solid", fgColor="A6A6A6"),
        "duty": PatternFill("solid", fgColor="A6A6A6"), "wish": PatternFill("solid", fgColor="EDEDED"),
        "break": PatternFill("solid", fgColor="F2F2F2")}
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def cell_text(c, kind):
    if not c:
        return "" if kind == "teacher" else "導師"
    t = c["type"]
    if kind == "class":
        if t == "subject":
            room = "" if c["room"] == "原班教室" else f"\n@{c['room'][:2]}"
            return f"{c['label']}\n{c['who']}{room}"
        if t == "sped":
            return f"{c['label']}（保留）\n抽離・導師上"
        if t == "keep":
            return f"導師（保留）\n{c.get('note', '')}"
        if t == "lock":
            return f"{c['label']}（鎖定）\n{c.get('who', '')}"
        if t == "lang":
            return f"{c['label']}\n{c.get('who', '')}"
        return c["label"]
    if t == "subject":
        room = "原班" if c["room"] == "原班教室" else c["room"][:2]
        return f"{c['label']}\n{c['subject']}・{room}"
    return c["label"]


def grid_block(ws, top, left, title, cells, kind, footer=""):
    ws.cell(top, left, title).font = Font(bold=True, size=12)
    ws.merge_cells(start_row=top, start_column=left, end_row=top, end_column=left + 5)
    for j, h in enumerate(["節次", "一", "二", "三", "四", "五"]):
        c = ws.cell(top + 1, left + j, h)
        c.fill, c.border, c.alignment, c.font = FILL["head"], BOX, CENTER, Font(bold=True)
    for i, p in enumerate(ROWS):
        rr = top + 2 + i
        head = ws.cell(rr, left, f"第{p}節" if isinstance(p, int) else p)
        head.fill, head.border, head.alignment = FILL["head"], BOX, CENTER
        for j, dd in enumerate("一二三四五"):
            c = ws.cell(rr, left + 1 + j)
            c.border, c.alignment = BOX, CENTER
            if not isinstance(p, int):
                c.fill = FILL["break"]
                continue
            v = cells.get(f"{dd}{p}")
            c.value = cell_text(v, kind)
            if v and v["type"] in FILL:
                c.fill = FILL[v["type"]]
                if v["type"] == "off":
                    c.font = Font(color="FFFFFF")
            if v and v["type"] == "subject":
                c.font = Font(bold=True)
            if not v and kind == "class":
                c.font = Font(color="808080", italic=True)
        ws.row_dimensions[rr].height = 34 if isinstance(p, int) else 12
    if footer:
        ws.cell(top + 11, left, footer).font = Font(size=9, color="404040")


def grid_sheet(wb, name, items, kind):
    ws = wb.create_sheet(name)
    for col in range(1, 14):
        ws.column_dimensions[get_column_letter(col)].width = 6 if col in (1, 8) else (2 if col == 7 else 11.5)
    for n, (title, cells, footer) in enumerate(items):
        block_row, side = divmod(n, 2)
        top = 1 + block_row * 13
        grid_block(ws, top, 1 + side * 7, title, cells, kind, footer)
        if side == 1 and block_row % 3 == 2:
            ws.row_breaks.append(Break(id=top + 12))
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4


def write_xlsx(r, out):
    wb = Workbook()
    ws = wb.active
    ws.title = "規則體檢"
    m = r["meta"]
    ws["A1"] = "115 學年度 科任排課結果　規則體檢"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"資料來源：{m['source']}　產生時間：{m['generated']}　結果：{m['status']}（{m['seconds']} 秒）"
    heads = ["編號", "層級", "規則積木", "對象", "內容", "結果", "說明"]
    for j, h in enumerate(heads):
        c = ws.cell(4, 1 + j, h)
        c.fill, c.border, c.font = FILL["head"], BOX, Font(bold=True)
    for i, rule in enumerate(r["rules"]):
        mark = {True: "○ 做到", False: "△ 部分", None: "— 略過"}[rule["ok"]]
        vals = [rule["id"], rule["level"], rule["block"], rule["target"], rule["content"], mark, rule["summary"]]
        for j, v in enumerate(vals):
            c = ws.cell(5 + i, 1 + j, v)
            c.border = BOX
            c.alignment = Alignment(vertical="center", wrap_text=True)
    row = 6 + len(r["rules"])
    for item in r.get("insights", []) + r.get("issues", []):
        ws.cell(row, 1, "提醒").font = Font(bold=True)
        ws.cell(row, 2, item["text"])
        row += 1
    for col, w in zip("ABCDEFG", [6, 7, 13, 14, 36, 9, 40]):
        ws.column_dimensions[col].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    classes = [(f"{c}　（{v['building']} {v['room']}）", v["cells"],
                f"科任 {v['subject_hours']} 節｜保留給導師 {v['keep']} 節｜導師自由排 {v['free']} 格")
               for c, v in r["classes"].items()]
    grid_sheet(wb, "班級課表", classes, "class")
    teachers = [(f"{t}　（{'、'.join(v['subjects'])}）", v["cells"], f"每週 {v['hours']} 節｜每天 {v['per_day']}")
                for t, v in sorted(r["teachers"].items())]
    grid_sheet(wb, "老師課表", teachers, "teacher")

    ws = wb.create_sheet("全校總表")
    slots = [f"{d}{p}" for d in r["days"] for p in r["periods"]]
    ws.cell(1, 1, "班級").fill = FILL["head"]
    for j, k in enumerate(slots):
        c = ws.cell(1, 2 + j, f"{k[0]}\n{k[1]}")
        c.fill, c.alignment, c.border = FILL["head"], CENTER, BOX
        ws.column_dimensions[get_column_letter(2 + j)].width = 4.2
    for i, (cname, v) in enumerate(r["classes"].items()):
        ws.cell(2 + i, 1, cname).border = BOX
        for j, k in enumerate(slots):
            cell = v["cells"].get(k)
            c = ws.cell(2 + i, 2 + j, (cell["label"][:2] if cell and cell["type"] != "off" else ""))
            c.border, c.alignment = BOX, CENTER
            c.font = Font(size=8, bold=bool(cell and cell["type"] == "subject"))
            if cell and cell["type"] in FILL:
                c.fill = FILL[cell["type"]]
    ws.freeze_panes = "B2"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(out)


def main():
    src = Path(sys.argv[1])
    r = json.loads(src.read_text(encoding="utf-8"))
    if "rules" in r:
        r["insights"] = insights(r)
        src.write_text(json.dumps(r, ensure_ascii=False, indent=1), encoding="utf-8")
        write_xlsx(r, src.parent / "排課結果_灰階列印.xlsx")
    write_html(r, src.parent / "課表檢視.html")
    print("已輸出：課表檢視.html、排課結果_灰階列印.xlsx")


if __name__ == "__main__":
    main()
